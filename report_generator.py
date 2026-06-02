import os
import requests
import time
import json
import re
import pandas as pd
import threading
import tkinter as tk
from tkinter import messagebox
import customtkinter as cctk
from tkcalendar import Calendar
from datetime import timedelta, datetime, timezone, time as dt_time
import pytz
from openpyxl.styles import PatternFill, Font, Alignment
from concurrent.futures import ThreadPoolExecutor, as_completed
from html import unescape
import traceback
# ---------------------------
# CONFIG: GHL Cuentas
# ---------------------------
ACCOUNTS = [
    {
        "name": "R2.1",
        "location_id": "gPO6FiXFjyul3qojXQKh",
        "stage_id": "dea2dacb-97d4-4165-b95b-5a02f6a87155",
        "custom_field": "fydz3DjQN2pib2axnFZe",
        "dataventa_id": "7veOYT8o670WKCIbRd92",
        "token": "pit-aa32ae57-f021-4345-9c17-599788ca222d",
        "secuencia_cf": "CDP7RyHYQIhVRNtmuYOX",
        "anuncio_cf": "EsjAWJ6IbHmCB47SbGTe",
        "primer_mensaje_cf": "ILBHzfqHvQpVXZqOqPTt"
    },
    {
        "name": "R2.2",
        "location_id": "HK2q1x0usZEQbnuJCD6B",
        "stage_id": "a5b50462-e82e-4256-b83d-d613ad20abcd",
        "custom_field": "U6OdqsQDVQTUDxxNBtvp",
        "dataventa_id": "1L7km8XXq3V1kDMRyvCC",
        "token": "pit-c8f89986-1a34-4178-9064-6f678a697a01",
        "secuencia_cf": "deDXbK9Nw5khnuk679CM",
        "anuncio_cf": "vzIYG6B7mypHxXM9kLGi",
        "primer_mensaje_cf": "gMCJJq0vMZYN0cz1BnUj"
    },
    {
        "name": "R2.3",
        "location_id": "jXN4id73HVqpa75YOR1N",
        "stage_id": "d94817e9-a7fb-4ea3-bed0-d1f117001825",
        "custom_field": "xftnXlHb41aDvIx8N26d",
        "dataventa_id": "JJYIQaDMKpMMekORmcqy",
        "token": "pit-cfd6cb24-6fee-441d-b01f-415160eb6f7b",
        "secuencia_cf": "ENXxDMYOkFK7XnjzrdOU",
        "anuncio_cf": "r8sOhHm5PKNtXv65jSaN",
        "primer_mensaje_cf": "vphhYZkFfAiaSrgfTiA3"
    },
    {
        "name": "R1.1",
        "location_id": "GmvsWG2a09UJjYwzFwN7",
        "stage_id": "c8d3a128-f03d-4bce-b61d-50593b7c4ebc",
        "custom_field": "k6xXzN1ksd16i1n68o4P",
        "dataventa_id": "kHdjdcIduLvTr3nL6DgZ",
        "token": "pit-cc6a4560-665e-4063-b12e-7dd3ec412570",
        "secuencia_cf": "rG3qADV2DReag3vE4LnZ",
        "anuncio_cf": "hN5VqUs4cgWP4tNEdnTI",
        "primer_mensaje_cf": "at9WccT1xJ4tJRm5KpbK"
    },
    {
        "name": "R1.2",
        "location_id": "9rHHeTsNpfJuiUkOoLdM",
        "stage_id": "59f6eff0-f06b-4f5a-b2ae-21f50ec8af32",
        "custom_field": "o7giXoy1LK8KMuzH2FNi",
        "dataventa_id": "DeNNFP4LihWoLaIpG0B2",
        "token": "pit-51105ced-165a-437d-bf76-37c9ec75f00e",
        "secuencia_cf": "AAYTXtJX7jRHPn0VDqVH",
        "anuncio_cf": "IELF1xRsnl1nWvoHBRHY",
        "primer_mensaje_cf": "uEKcNGLJvv7znfVJb2Z4"
    },
    {
        "name": "R1.3",
        "location_id": "riT0De9iiwhd84gSRco3",
        "stage_id": "3b346e40-01ea-416e-98ff-60ee147aded1",
        "custom_field": "f310g4Z4A1OxHl3KYLRv",
        "dataventa_id": "4tDKaHgodZXQ4ewzjJyh",
        "token": "pit-58f93f09-e7a8-40d0-8ce6-32bf2f2b87b4",
        "secuencia_cf": "7s52aRnEuz8T3v0H3IlN",
        "anuncio_cf": "Vki3QtqsNJtYcRC4Fkzt",
        "primer_mensaje_cf": "78eH2yy4INV88QguossT"
    }
]

VENDEDOR_MAP = {
    "MARIA RENE SANTA CRUZ COSAJAY": "MARIA SANTACRUZ",
    "HENRY ESTUARDO PACHECO ARIANO": "HENRY PACHECO",
    "ROSA LIDIA PEREZ": "ROSA PEREZ",
    "CLEMENCIA ROCIO MIZA": "CLEMENCIA MIZA",
    "GRECIA SARAI FLORES FLORES": "GRECIA FLORES",
    "EVELYN PAOLA DARDON MORATAYA": "EVELYN DARDON",
    "MARVIN FRANCISCO CARRERA PINEDA": "MARVIN CARRERA",
    "YEIMI NOHEMI HERNANDEZ GOMEZ": "YEIMI HERNANDEZ",
    "ANA KARINA VELASQUEZ CALDERON DE CANAS": "ANA VELASQUEZ",
    "JHONATAN ISRAEL BARRIOS MUÑOZ": "JHONATAN BARRIOS",
    "CRISTIAN OMAR SANTOS ROSALES": "Cristian Santos",
    "LOURDES CAROLINA CAMPOS REYES": "Lourdes Campos",
    "CARLOS ARMANDO GIL TAJIN": "Carlos Gil",
    "OSCAR DANIEL IXCAYAU AGUILAR": "Oscar Ixcayau",
    "SIN ASIGNAR": "SinAsignar",
    "BYRON SILVERIO ORTIZ MENDOZA": "BYRON ORTIZ",
    "ESTHER DEL CARMEN LOPEZ CRUZ": "ESTHER LOPEZ",
    "SONIA MARIBEL CHIROY": "SONIA CHIROY",
    "MARVIN DAVID CANEL HERNANDEZ": "MARVIN CANEL",
    "BLANCA  RUTILIA YACAB AC": "BLANCA YACAB",
    "PEDRO JOSUE YAX VILLATORO": "PEDRO VILLATORO",
    "VICTOR MIGUEL ANGEL ASTURIAS TZAMOL": "VICTOR ANGEL",
    "ERICKA VICTORIA GONZALEZ LÓPEZ": "ERICKA GONZALEZ",
    "KAREN YULISA MENCOS RODAS": "KAREN MENCOS",
    "FABIANA MENDOZA LOPEZ": "FABIANA LOPEZ",
    "STEPHANIE DENNES AJUCHAN HERNANDEZ": "Stephanie Ajuchan",
    "ALISON MELISA AJANEL JOLON": "Alyson Ajanel",
    "IRMA LOURDES DE PAZ": "Lourdes Paz",
    "DENIS SEBASTIAN AJUCHAN SANTOS": "Denis Ajuchan",
    "SANDRA PATRICIA GUTIERREZ BLANCO": "Sandra Gutierrez",
    "FELIX  ANTONIO ACEITUNO BARRIENTOS": "FELIX ACEITUNO",
    "JUAN LUIS COTZAJAY TIJE": "JUAN COTZAJAY",
    "BRIAN SALVADOR CANO RAMIREZ": "BRIAN CANO",
    "JAMNIA GALILEA MORALES ORELLANA": "GALILEA MORALES",
    "WILMER ALEXANDER CAJAS JUAREZ": "WILMER CAJAS",
    "ADA LUZ MARIA MENDEZ SOSA": "Ada luz Mendez",
    "CHRISTIAN OLIVER GONZALEZ ORDOÑEZ": "Christian Gonzalez",
    "JESSICA ILEANA GARCIA PALACIOS": "JESSICA GARCIA",
    "GERBER ROMUALDO YAX VILLATORO": "GERBER YAX",
    "HAZEL BARINIA AGUILAR GONZALEZ": "HAZEL AGUILAR",
    "OSCAR JOSUE HERNANDEZ PEREZ": "Oscar Hernandez",
    "WALTER NEHEMIAS GUERRA": "WALTER GUERRA",
    "YESSICA ALEJANDRA CARRERA PINEDA": "YESSICA CARRERA",
    "HIDELKY AJIN": "HidelkyAjin",
    "BYRON ALEXIS MAYOR": "BYRON MAYOR",
    "YARELIN BARRAZA ARIAS": "YARELIN ARIAS",
    "YOSELIN EUFEMIA BARRAZA ARIAS": "YOSELIN BARRAZA",
    "YENDY MIREYA CUMAR CASTRO": "YENDY CUMAR",
    "ANGEL DANIEL LOPEZ PATZAN": "Angel Lopez",
    "BRYAN ARMANDO LOPEZ PATZAN": "Bryan Lopez",
    "HILLARY GISSELL RUCAL GALLINA": "Hillary Rucal",
    "KIMBERLY JESSENIA REGUAN RABAY": "Kimberly Reguán",
    "LUIS DAVID QUEXEL GIL": "Luis Quexel",
    "WENDY ARELIS ALQUIJAY GIL": "Wendy Alquijay",
    "CARLOS GIL TAJIN": "Carlos Gil",
    "ROSALINDA EUSEBIA RAMIREZ TEZEN": "Rosalinda Ramirez",
    "JULIO ALEJANDRO AJXUP GIL": "Julio Ajxup",
    "JOCARI ANASOL LOPEZ SICAL": "Jocari Lopez",
    "DIEGO SANTA CRUZ": "DIEGO SANTACRUZ",
    "ODILIA NINETTE CALEL CARAU": "ODILIA NINETH CALEL",
}

API_VERSION_OPPS = "2023-02-21"
API_VERSION_CONTACTS = "2021-07-28"
GUATEMALA_TZ = pytz.timezone("America/Guatemala")

# ---------------------------
# CONFIG: Facebook Ads
# ---------------------------
FB_ACCESS_TOKEN = "EAAQlsSqsOJkBQ8ELEZCLxm0CPEiqaUSoYw0oHB7ML7xvufZBn2B6t1bizlxBtv8gjc1r4bHiqlV0AHtbI9FDLTRhivFwpDr2xMzk7Waj8htSHanBW63gZCOPwPoVOgZBuNurP2ZB6vegJGxYRAmIR2Wp2JbPZAn9u4CVdHZC06TzaIngYeZAh3n4iaMtw3SyZA7XsNAZDZD"
FB_API_VERSION = "v19.0"
FB_AD_ACCOUNTS = ["act_622689460111355", "act_934171589566820"]
FB_USD_ACCOUNTS = ["act_934171589566820"]
USD_TO_GTQ = 7.8
FB_BASE_URL = f"https://graph.facebook.com/{FB_API_VERSION}"
# Se agregó account_id para aplicar la conversión a GTQ según la cuenta
FB_FIELDS_INSIGHTS = "account_id,ad_id,ad_name,adset_name,campaign_name,impressions,spend,clicks,actions,date_start,date_stop"

# ---------------------------
# CONFIG: Archivos Locales
# ---------------------------
PATH_METAS = "metas.xlsx"

# Regex para extracción
ANUNCIO_REGEX = re.compile(r"([A-Z]\d{3,4}[A-Z]\d{3})", re.IGNORECASE)
SECUENCIA_REGEX = re.compile(r"([A-Z]\d\.\d)", re.IGNORECASE)

# ---------------------------
# Backend Functions: General
# ---------------------------
def clean_html(raw_html):
    if not raw_html: return ""
    return unescape(re.sub(r'<[^>]+>', '', raw_html)).strip()

def format_date_ghl(val):
    """Retorna DD/MM/YYYY para Excel."""
    if not val: return ""
    if isinstance(val, (int, float)):
        try: return datetime.fromtimestamp(val / 1000, tz=timezone.utc).strftime("%d/%m/%Y")
        except: return str(val)
    s = str(val).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
    if len(s) >= 10 and s[2] == "/" and s[5] == "/":
        return s[:10]
    return s

def get_yyyy_mm_dd(val):
    """Retorna YYYY-MM-DD para comparación."""
    if not val: return ""
    if isinstance(val, (int, float)):
        try: return datetime.fromtimestamp(val / 1000, tz=timezone.utc).strftime("%Y-%m-%d")
        except: return ""
    s = str(val).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    if len(s) >= 10 and s[2] == "/" and s[5] == "/":
        return f"{s[6:10]}-{s[3:5]}-{s[0:2]}"
    return ""

def calculate_nit(nit, tel1, ghl_phone):
    n, t1, gp = str(nit).strip().upper(), str(tel1).strip(), str(ghl_phone).strip()
    res = n
    if n in ("CF", "C/F", ""): res = t1
    if not res:
        if len(gp) >= 12: res = gp[4:12]
        elif len(gp) > 4: res = gp[4:]
    return res

def get_custom_value(field):
    if not field or not isinstance(field, dict): return ""
    if "fieldValueDate" in field and field["fieldValueDate"]: return format_date_ghl(field["fieldValueDate"])
    if "fieldValueString" in field and field["fieldValueString"]: return str(field["fieldValueString"])
    if "fieldValue" in field and field["fieldValue"] is not None:
        v = field["fieldValue"]
        if isinstance(v, (int, float)) and v > 1000000000000: return format_date_ghl(v)
        return str(v)
    if "value" in field and field["value"] is not None:
        v = field["value"]
        if isinstance(v, list): return ", ".join(map(str, v))
        return str(v)
    return ""

def get_custom_fields_map(location_id, token):
    url = f"https://services.leadconnectorhq.com/locations/{location_id}/customFields?model=opportunity"
    headers = {"Authorization": f"Bearer {token}", "Version": API_VERSION_OPPS, "Accept": "application/json"}
    try:
        r = requests.get(url, headers=headers, timeout=30)
        if r.status_code != 200: return {}
        return {f.get("id"): f.get("name") for f in r.json().get("customFields", []) if isinstance(f, dict)}
    except: return {}

def get_users_by_location(location_id, token, version=API_VERSION_OPPS):
    url = f"https://services.leadconnectorhq.com/users/?locationId={location_id}"
    headers = {"Authorization": f"Bearer {token}", "Version": version, "Accept": "application/json"}
    try:
        r = requests.get(url, headers=headers, timeout=30)
        if r.status_code != 200: return {}
        return {u.get("id"): f"{u.get('firstName','') or ''} {u.get('lastName','') or ''}".strip() or u.get("email", "Desconocido") for u in r.json().get("users", [])}
    except: return {}

def safe_post(url, token, payload, version):
    headers = {"Authorization": f"Bearer {token}", "Version": version, "Content-Type": "application/json"}
    for attempt in range(1, 6):
        try:
            r = requests.post(url, headers=headers, json=payload, timeout=30)
            if r.status_code in (200, 201): return r.json()
            if r.status_code == 429:
                time.sleep(attempt * 2)
                continue
            return {"__error_status": r.status_code, "__error_text": r.text}
        except:
            time.sleep(attempt * 1.5)
            continue
    return {}

def parse_dataventa(dv_str):
    try:
        data = json.loads(dv_str)
        nit, dep, mun = data.get("nit", ""), data.get("departamento", ""), data.get("municipio", "")
        t1, t2, fv, nom = data.get("tel1", ""), data.get("tel2", ""), data.get("fechaVenta", ""), data.get("nombre", "")
        p_cols = {}
        all_slots = ["Camas y Combos SKU", "Cantidad Camas y Combo SKU", "Camas y Combos SKU1", "Cantidad Camas y Combo SKU1", "Cocinas SKU", "Cantidad Cocinas SKU", "Cocinas SKU1", "Cantidad Cocinas SKU1", "Salas SKU", "Cantidad Salas SKU", "Salas SKU1", "Cantidad Salas SKU1"]
        for i, p in enumerate(data.get("productos", [])):
            if i * 2 < len(all_slots): p_cols[all_slots[i*2]] = p.get("sku", ""); p_cols[all_slots[i*2 + 1]] = p.get("cantidad", "")
            else: p_cols[f"Producto_{i+1} SKU"] = p.get("sku", ""); p_cols[f"Cantidad Producto_{i+1}"] = p.get("cantidad", "")
        return nit, dep, mun, t1, t2, fv, nom, p_cols
    except: return "", "", "", "", "", "", "", {}

def get_mapped_vendedor(raw_vendedor):
    if not raw_vendedor: return "SinAsignar"
    nm = str(raw_vendedor).strip().upper()
    # Priorizar búsqueda exacta en el mapa
    if nm in VENDEDOR_MAP: return VENDEDOR_MAP[nm]
    # Buscar si alguna llave está contenida o viceversa (más flexible)
    for k, v in VENDEDOR_MAP.items():
        if k in nm or nm in k: return v
    return nm

def parse_ventas_unnested(dv_str, contact_id, opp_id, ghl_phone, vendedor, ghl_name="", sale_date_str="", secuencia=""):
    data = {}
    if dv_str:
        try: data = json.loads(dv_str)
        except: pass
    nit_j, t1 = data.get("nit", ""), data.get("tel1", "")
    try: total_docto = float(str(data.get("Total_General", 0)).replace(',', ''))
    except: total_docto = 0.0
    vendedor_final = get_mapped_vendedor(vendedor)
    base_metadata = {"ID CONTACTO": contact_id, "NIT": calculate_nit(nit_j, t1, ghl_phone), "NOMBRE": data.get("nombre", ghl_name), "TEL1": t1, "TEL2": data.get("tel2", ""), "VENDEDOR": vendedor_final, "MUNICIPIO": data.get("municipio", ""), "DIRECCION": data.get("direccion", ""), "RCF": "000000000000", "canal": data.get("canal", ""), "DEPARTAMENTO": data.get("departamento", ""), "FECHA": format_date_ghl(data.get("fechaVenta", sale_date_str)), "MARCA": data.get("marca", ""), "UBICACION": data.get("ubicacion", ""), "ANILLO": data.get("anillo", ""), "COMENTARIOS": data.get("COMENTARIO", ""), "SECUENCIA": secuencia, "ID Oportunidad": opp_id, "BODEGAF": str(data.get("bodega", ""))[:4] if data.get("bodega") else "", "TOTAL DOCTO": total_docto}
    empty_metadata = {k: "" for k in base_metadata.keys()}
    prods = data.get("productos", [])
    if not prods: return [base_metadata]
    rows = []
    for i, p in enumerate(prods):
        row = dict(base_metadata) if i == 0 else dict(empty_metadata)
        sku_f = str(p.get("sku", ""))
        parts = sku_f.split("_")
        try: precio_combo = float(str(p.get("precio", 0)).replace(',', ''))
        except: precio_combo = 0.0
        row.update({"SKU": parts[0] if parts else sku_f, "DESCRIPCION": parts[1] if len(parts) > 1 else "", "Cantidad de combo": p.get("cantidad", ""), "PRECIO COMBO": precio_combo}); rows.append(row)
    return rows

def make_utc_range(start_date, end_date):
    start_local = GUATEMALA_TZ.localize(datetime.combine(start_date, dt_time.min))
    end_local = GUATEMALA_TZ.localize(datetime.combine(end_date, dt_time.max))
    return start_local.astimezone(pytz.UTC).isoformat().replace("+00:00", "Z"), end_local.astimezone(pytz.UTC).isoformat().replace("+00:00", "Z")

def cargar_metas(filepath, log_callback=None):
    if not os.path.exists(filepath):
        if log_callback: log_callback(f"Advertencia: No se encontró {filepath}")
        return pd.DataFrame()
    try:
        df = pd.read_excel(filepath)
        if log_callback: log_callback(f"Metas cargadas: {len(df)} registros desde {filepath}")
        # Limpieza similar a Power Query
        for col in df.select_dtypes(include=['object', 'str']).columns:
            df[col] = df[col].astype(str).str.strip().str.upper().replace("NAN", "").replace("NONE", "")
        # Normalizar nombres de columnas (SUB ANILLO -> SUB_ANILLO para JS)
        df.columns = [c.replace(" ", "_").replace("/", "_").replace("%", "PORC") for c in df.columns]
        return df
    except Exception as e:
        print(f"Error cargando metas: {e}")
        return pd.DataFrame()

# ---------------------------
# Backend Functions: GHL
# ---------------------------
def extraer_datos_anuncio(text):
    if not text: return "", ""
    match_anu = ANUNCIO_REGEX.search(text)
    if not match_anu: return "", ""
    anuncio = match_anu.group(1)
    tipo_post = ""
    pattern_tipo = re.escape(anuncio) + r"\.([^.]+)"
    match_tipo = re.search(pattern_tipo, text, re.IGNORECASE)
    if match_tipo: tipo_post = match_tipo.group(1)
    return anuncio, tipo_post

def extraer_secuencia(text):
    if not text: return ""
    t = text.strip().upper()
    match_sec = SECUENCIA_REGEX.search(t)
    if match_sec: return match_sec.group(1)
    # Soporte para formato A03-A
    match_alt = re.search(r"([A-Z]\d{1,2}-[A-Z])", t)
    if match_alt: return match_alt.group(1)
    return t if len(t) <= 8 else "" # Si es corto lo tomamos como código

def fetch_contacts_for_account(acc, start_utc, end_utc, log_callback):
    token, loc, acc_name = acc["token"], acc["location_id"], acc["name"]
    sec_cf, anu_cf, pm_cf = acc["secuencia_cf"], acc["anuncio_cf"], acc["primer_mensaje_cf"]
    log_callback(f"Extraer Contactos: {acc_name}...")
    u_map = get_users_by_location(loc, token, version=API_VERSION_CONTACTS)
    all_contacts, page, limit = [], 1, 100
    url = "https://services.leadconnectorhq.com/contacts/search"
    while True:
        payload = {"locationId": loc, "page": page, "pageLimit": limit, "filters": [{"field": "dateAdded", "operator": "range", "value": {"gt": start_utc, "lt": end_utc}}]}
        res = safe_post(url, token, payload, API_VERSION_CONTACTS)
        if not res: break
        if isinstance(res, dict) and res.get("__error_status"):
            log_callback(f"  {acc_name} ERROR {res.get('__error_status')}: {res.get('__error_text')[:100]}")
            break
        contacts = res.get("contacts", [])
        if not isinstance(contacts, list) or not contacts: break
        all_contacts.extend(contacts); page += 1
        if len(contacts) < limit: break

    formatted_contacts = []
    for c in all_contacts:
        uid = c.get("assignedTo")
        assigned_name = get_mapped_vendedor(u_map.get(uid, "")) if uid else "SinAsignar"
        date_iso, date_fmt, dt_local = c.get("dateAdded"), "", None
        if date_iso:
            dt_local = datetime.fromisoformat(date_iso.replace("Z", "+00:00")).astimezone(GUATEMALA_TZ)
            date_fmt = f"{dt_local.day}/{dt_local.month:02d}/{dt_local.year}"
        secuencia_raw, anuncio_raw, primer_mensaje_texto = "", "", ""
        for cf in c.get("customFields", []):
            cid, val = cf.get("id"), get_custom_value(cf)
            if cid == sec_cf: secuencia_raw = val.strip()
            elif cid == anu_cf: anuncio_raw = val.strip()
            elif cid == pm_cf: primer_mensaje_texto = val.strip()
        anuncio, tipo_post = extraer_datos_anuncio(anuncio_raw)
        if not anuncio and primer_mensaje_texto: anuncio, tipo_post = extraer_datos_anuncio(primer_mensaje_texto)
        secuencia = acc_name # Usar nombre de cuenta estrictamente
        formatted_contacts.append({
            "id": c.get("id", ""),
            "fecha": date_fmt,
            "fecha_iso": dt_local.strftime("%Y-%m-%d") if dt_local else "",
            "asignado": assigned_name,
            "secuencia": acc_name,
            "anuncio": anuncio,
            "tipo_post": tipo_post
        })
    return formatted_contacts

def fetch_for_account(acc, ghl_start, ghl_end, client_start, client_end, log_callback):
    token, loc, stage, cfield, dv_id, acc_name = acc["token"], acc["location_id"], acc["stage_id"], acc["custom_field"], acc["dataventa_id"], acc["name"]
    log_callback(f"Extraer Ventas: {acc_name}...")
    u_map, cf_names, all_opps, page, limit = get_users_by_location(loc, token), get_custom_fields_map(loc, token), [], 1, 100
    url = "https://services.leadconnectorhq.com/opportunities/search"
    while True:
        # Volviendo al formato original robusto de filtros anidados
        payload = {
            "locationId": loc, "page": page, "limit": limit,
            "filters": [{"group": "AND", "filters": [
                {"field": "pipeline_stage_id", "operator": "eq", "value": stage},
                {"field": "status", "operator": "eq", "value": "won"}
            ]}],
            "sort": [{"field": "date_added", "direction": "desc"}],
            "additionalDetails": {"notes": True}
        }
        res = safe_post(url, token, payload, API_VERSION_OPPS)
        if not res or (isinstance(res, dict) and res.get("__error_status")):
            if isinstance(res, dict) and res.get("__error_status"):
                log_callback(f"  {acc_name} ERROR {res.get('__error_status')}: {res.get('__error_text')[:100]}")
            break

        opps = res.get("opportunities", [])
        if not isinstance(opps, list) or not opps: break
        all_opps.extend(opps); page += 1
        if len(opps) < limit: break
        if page > 100: break

    r_opps, r_ventas, filtered_count = [], [], 0
    log_callback(f"  {acc_name}: {len(all_opps)} ganadas encontradas en total. Filtrando por fecha...")
    for op in all_opps:
        if not isinstance(op, dict): continue
        try:
            opp_cfs = op.get("customFields") or op.get("custom_fields") or []
            sale_date_iso, sale_date_str, dv_str, cf_data = "", "", "", {}
            for cf in opp_cfs:
                fid, val = cf.get("id"), get_custom_value(cf)
                if fid == cfield:
                    sale_date_str, sale_date_iso = val, get_yyyy_mm_dd(cf.get("fieldValueDate") or cf.get("fieldValue") or cf.get("fieldValueString"))
                if fid == dv_id: dv_str = str(cf.get("fieldValue") or cf.get("fieldValueString") or "")
                fname = cf_names.get(fid, fid)
                if fname and fname.strip().lower() != "id de oportunidad": cf_data[fname] = val
            if not sale_date_iso:
                # Fallback: intentar usar updatedAt si la fecha custom no está
                updated_at_iso = get_yyyy_mm_dd(op.get("updatedAt") or op.get("updated_at"))
                if updated_at_iso:
                    sale_date_iso = updated_at_iso

            if not sale_date_iso or not (client_start <= sale_date_iso <= client_end):
                filtered_count += 1
                continue
            vendedor_raw, gnam, opp_id_val = get_mapped_vendedor(u_map.get(op.get("assignedTo") or op.get("assigned_to"), "")), (op.get("contact", {}).get("name", "") if isinstance(op.get("contact"), dict) else ""), op.get("id", "")
            dv_data = json.loads(dv_str) if dv_str else {}
            anu_val, _ = extraer_datos_anuncio(dv_data.get("anuncio", "") or dv_data.get("Anuncio", "")); row = {"Asignado": vendedor_raw, "Secuencia": acc_name, "Anuncio": anu_val, "fecha_iso": sale_date_iso, "Mes": int(sale_date_iso[5:7]) if sale_date_iso else 0, "Anio": int(sale_date_iso[0:4]) if sale_date_iso else 0, "fase": op.get("pipelineStageName", "Cierre de Venta"), "Valor del cliente potencial": op.get("monetaryValue", 0), "asignado": vendedor_raw, "Creado": format_date_ghl(op.get("createdAt")), "Ultimo Actualizado": format_date_ghl(op.get("updatedAt")), "Seguidores": "", "Notas": " | ".join([clean_html(n.get("body", "")) for n in op.get("notes", []) if isinstance(n, dict)]), "etiquetas": ", ".join(op.get("tags", [])) if isinstance(op.get("tags"), list) else "", "estado": op.get("status", ""), "ID de contacto": op.get("contactId", ""), "Cliente": gnam, "Cod": str(opp_id_val)[:10], "MARCA": dv_data.get("marca", ""), "ANILLO": dv_data.get("anillo", ""), "UBICACION": dv_data.get("ubicacion", ""), "Mes": int(sale_date_iso[5:7]) if sale_date_iso else "", "DataVenta": dv_str, "ID de oportunidad": opp_id_val}
            row.update(cf_data)
            nit_j, dep, mun, t1, t2, fv_j, nom_j, p_cols = parse_dataventa(dv_str)
            gp, f_final = (op.get("contact", {}).get("phone", "") if isinstance(op.get("contact"), dict) else ""), format_date_ghl(sale_date_str or fv_j)
            row.update({"NIT": calculate_nit(nit_j, t1, gp), "Departamento": dep, "Municipio": mun, "Telefono 1": t1, "Telefono 2": t2, "Fecha": f_final, "Fecha de Venta": f_final})
            if nom_j: row["Cliente"] = nom_j
            row.update(p_cols); r_opps.append(row); r_ventas.extend(parse_ventas_unnested(dv_str, op.get("contactId", ""), op.get("id", ""), gp, vendedor_raw, gnam, sale_date_str, acc_name))
        except: continue
    log_callback(f"  {acc_name}: {len(r_opps)} ventas aceptadas, {filtered_count} fuera de rango.")
    return r_opps, r_ventas

# ---------------------------
# Backend Functions: Facebook
# ---------------------------
def extraer_precio_fb(nombre):
    if not isinstance(nombre, str): return None
    m = re.search(r'\.(\d+)', nombre)
    return int(m.group(1)) if m else None

def fb_api_get(url, params):
    params["access_token"] = FB_ACCESS_TOKEN
    try: r = requests.get(url, params=params, timeout=30); return r.json()
    except: return {}

def mapping_secuencia_gasto(camp):
    camp = str(camp).upper()
    if camp.startswith("DIEGOA01C1.PAGINA M"): return "R1.3"
    if camp.startswith("DIEGO"): return "R1.3"
    if camp.startswith("VICTORIAA02C1"): return "R1.2"
    if camp.startswith("VICTORIA"): return "R1.2"
    if camp.startswith("NOHEA02C1"): return "R1.2"
    if camp.startswith("NOHEA02C2"): return "R1.3"
    if camp.startswith("NOHE"): return "R1.3"
    if camp.startswith("ANGEL"): return "R1.3"
    if camp.startswith("2510"): return "R1.1"
    if camp.startswith("RRHH"): return "RRHH"
    if camp.startswith("TIENDAS"): return "R1.1"
    if camp.startswith("BOT2"): return "R1.2"
    if "R2.2" in camp: return "R2.2"
    if "R2.1" in camp: return "R2.1"
    if "R1.1" in camp: return "R1.1"
    if "R1.2" in camp: return "R1.2"
    if "R1.3" in camp: return "R1.3"
    if "R2.3" in camp: return "R2.3"
    if "R3.2" in camp: return "R3.2"
    # Fallback to ACCOUNT names check
    for r_name in ["R1.1", "R1.2", "R1.3", "R2.1", "R2.2", "R2.3", "R3.2"]:
        if r_name in camp: return r_name
    return "OTRO"

def obtener_insights(account, fecha_desde, fecha_hasta, log_callback):
    log_callback(f"Extraer Insights FB: {account}...")
    url, params = f"{FB_BASE_URL}/{account}/insights", {"fields": FB_FIELDS_INSIGHTS, "level": "ad", "limit": 500, "time_range": json.dumps({"since": fecha_desde, "until": fecha_hasta}), "time_increment": 1, "filtering": '[{"field":"spend","operator":"GREATER_THAN","value":0}]'}
    data = []
    while True:
        js = fb_api_get(url, params)
        if "data" not in js: break
        data.extend(js["data"])
        if "paging" in js and "next" in js["paging"]: url, params = js["paging"]["next"], {}
        else: break
    return data

def obtener_creatives(ad_ids):
    creative_map = {}
    for i in range(0, len(ad_ids), 50):
        block = ad_ids[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{aid}?fields=creative"} for aid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list): continue
        for resp in responses:
            if resp.get("code") != 200: continue
            body = json.loads(resp.get("body", "{}"))
            ad_id, creative_id = body.get("id"), body.get("creative", {}).get("id")
            if ad_id and creative_id: creative_map[str(ad_id)] = str(creative_id)
        time.sleep(0.3)
    return creative_map

def obtener_paginas(creatives):
    page_map = {}
    if not creatives: return page_map
    for i in range(0, len(creatives), 50):
        block = creatives[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{cid}?fields=object_story_spec,actor_id,effective_object_story_id"} for cid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list): continue
        for resp in responses:
            if resp.get("code") != 200: continue
            body = json.loads(resp.get("body", "{}"))
            cid, oss, actor_id, post_id = body.get("id"), body.get("object_story_spec", {}), body.get("actor_id"), body.get("effective_object_story_id")
            page_id = oss.get("page_id") or oss.get("page", {}).get("id") or oss.get("link_data", {}).get("page_id") or oss.get("video_data", {}).get("page_id")
            if not page_id and post_id and "_" in post_id: page_id = post_id.split("_")[0]
            final_p_id = str(page_id) if page_id else (str(actor_id) if actor_id else None)
            if final_p_id: page_map[str(cid)] = final_p_id
        time.sleep(0.3)
    return page_map

def obtener_nombres_paginas(page_ids):
    names, page_ids = {}, [pid for pid in page_ids if pid]
    if not page_ids: return names
    for i in range(0, len(page_ids), 50):
        block = page_ids[i:i+50]
        batch = [{"method": "GET", "relative_url": f"{pid}?fields=name,username,about"} for pid in block]
        r = requests.post(f"{FB_BASE_URL}/", data={"access_token": FB_ACCESS_TOKEN, "batch": json.dumps(batch)})
        responses = r.json()
        if not isinstance(responses, list): continue
        for resp in responses:
            if resp.get("code") != 200: continue
            body = json.loads(resp.get("body", "{}"))
            pid, name = str(body.get("id")) if body.get("id") else None, body.get("name") or body.get("username") or body.get("about") or "Sin Nombre"
            if pid: names[pid] = name
        time.sleep(0.3)
    return names

def obtener_paginas_autorizadas():
    names, url, params = {}, f"{FB_BASE_URL}/me/accounts", {"limit": 100}
    while True:
        js = fb_api_get(url, params)
        if "data" not in js: break
        for item in js["data"]:
            pid, pname = str(item.get("id")), item.get("name")
            if pid and pname: names[pid] = pname
        if "paging" in js and "next" in js["paging"]: url, params = js["paging"]["next"], {}
        else: break
    return names

# ---------------------------
# DATE PICKER
# ---------------------------
class FloatingRangePicker(cctk.CTkFrame):
    def __init__(self, parent, title):
        super().__init__(parent, corner_radius=12); self.start_date, self.end_date, self.pop = None, None, None
        cctk.CTkLabel(self, text=title, font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=15, pady=(10,2))
        self.entry_frame = cctk.CTkFrame(self, fg_color="transparent"); self.entry_frame.pack(padx=15, pady=(2,10), fill="x")
        self.entry_start = cctk.CTkEntry(self.entry_frame, placeholder_text="Inicio", height=32, corner_radius=8, font=("Segoe UI", 11), justify="center", state="readonly"); self.entry_start.pack(side="left", expand=True, fill="x", padx=(0,2)); self.entry_start.bind("<Button-1>", lambda e: self.open_calendar(self.entry_start, "start"))
        cctk.CTkLabel(self.entry_frame, text="→", font=("Segoe UI", 12, "bold")).pack(side="left")
        self.entry_end = cctk.CTkEntry(self.entry_frame, placeholder_text="Fin", height=32, corner_radius=8, font=("Segoe UI", 11), justify="center", state="readonly"); self.entry_end.pack(side="left", expand=True, fill="x", padx=(2,0)); self.entry_end.bind("<Button-1>", lambda e: self.open_calendar(self.entry_end, "end"))

    def open_calendar(self, target_entry, date_type):
        if self.pop: self.close_calendar()
        self.update_idletasks(); self.pop = tk.Toplevel(self); self.pop.overrideredirect(True); self.pop.attributes("-topmost", True); self.pop.geometry(f"300x320+{target_entry.winfo_rootx()}+{target_entry.winfo_rooty() + 35}"); self.pop.grab_set()
        container = cctk.CTkFrame(self.pop, corner_radius=10, border_width=2, border_color="#76933C", fg_color="#ffffff"); container.pack(fill="both", expand=True)
        header = tk.Frame(container, bg="#ffffff", height=30); header.pack(fill="x", padx=8, pady=2)
        tk.Label(header, text="ELEGIR FECHA", font=("Segoe UI", 9, "bold"), fg="#333333", bg="#ffffff").pack(side="left")
        tk.Button(header, text="✕", font=("Arial", 9), bd=0, bg="#ffffff", command=self.close_calendar).pack(side="right")
        self.cal = Calendar(container, selectmode="day", date_pattern="yyyy-mm-dd", background='white', foreground='black', selectbackground='#1890ff', selectforeground='white', borderwidth=0); self.cal.pack(pady=5, padx=10, fill="both", expand=True)
        self.cal.bind("<<CalendarSelected>>", lambda e: self._on_date_selected(target_entry, date_type)); self.pop.focus_set()

    def close_calendar(self):
        if self.pop: self.pop.grab_release(); self.pop.destroy(); self.pop = None

    def _on_date_selected(self, target_entry, date_type):
        date_obj = datetime.strptime(self.cal.get_date(), "%Y-%m-%d").date()
        if date_type == "start": self.start_date = date_obj
        else: self.end_date = date_obj
        target_entry.configure(state="normal"); target_entry.delete(0, "end"); target_entry.insert(0, str(date_obj)); target_entry.configure(state="readonly"); self.after(200, self.close_calendar)

# ---------------------------
# APP
# ---------------------------
class App(cctk.CTk):
    def __init__(self):
        super().__init__(); self.title("DUPAZA PRO"); self.geometry("380x700"); cctk.set_appearance_mode("dark"); cctk.set_default_color_theme("green")
        self.grid_columnconfigure(0, weight=1); self.grid_rowconfigure(1, weight=1)
        header = cctk.CTkFrame(self, height=45, corner_radius=0); header.grid(row=0, column=0, sticky="ew")
        cctk.CTkLabel(header, text="📊 DUPAZA PRO", font=("Segoe UI", 18, "bold")).pack(pady=8)
        body = cctk.CTkScrollableFrame(self, fg_color="transparent"); body.grid(row=1, column=0, sticky="nsew", padx=8, pady=8)
        body.grid_columnconfigure(0, weight=1)
        self.contacts_picker = FloatingRangePicker(body, "📇 CONTACTOS GHL"); self.contacts_picker.grid(row=0, column=0, pady=3, sticky="ew")
        self.sales_picker = FloatingRangePicker(body, "💰 VENTAS GHL"); self.sales_picker.grid(row=1, column=0, pady=3, sticky="ew")
        self.fb_picker = FloatingRangePicker(body, "🔵 FACEBOOK ADS"); self.fb_picker.grid(row=2, column=0, pady=3, sticky="ew")
        self.generate_btn = cctk.CTkButton(self, text="🚀 GENERAR EXCEL", height=38, font=("Segoe UI", 14, "bold"), corner_radius=10, command=self.start_process); self.generate_btn.grid(row=2, column=0, pady=8, padx=15, sticky="ew")
        logs_frame = cctk.CTkFrame(self, height=80, corner_radius=10); logs_frame.grid(row=3, column=0, sticky="ew", padx=10, pady=(0,10))
        self.console = cctk.CTkTextbox(logs_frame, height=60, font=("Consolas", 10)); self.console.pack(fill="both", expand=True, padx=5, pady=5); self.log("LISTO.")

    def log(self, txt):
        txt = str(txt)
        hour = datetime.now().strftime("%H:%M:%S")
        self.console.configure(state="normal"); self.console.insert("end", f"[{hour}] {txt}\n"); self.console.see("end"); self.console.configure(state="disabled")

    def start_process(self):
        if not all([self.sales_picker.start_date, self.sales_picker.end_date, self.contacts_picker.start_date, self.contacts_picker.end_date, self.fb_picker.start_date, self.fb_picker.end_date]): messagebox.showwarning("Atención", "Elija todos los rangos."); return
        self.generate_btn.configure(state="disabled", text="🚀 PROCESANDO..."); threading.Thread(target=self.execute_logic, daemon=True).start()

    def execute_logic(self):
        try:
            sd_opp, ed_opp, sd_con, ed_con = self.sales_picker.start_date, self.sales_picker.end_date, self.contacts_picker.start_date, self.contacts_picker.end_date
            s_iso_o, e_iso_o, ghl_s_o, ghl_e_o = sd_opp.strftime("%Y-%m-%d"), ed_opp.strftime("%Y-%m-%d"), sd_opp.strftime("%Y-%m-%dT00:00:00.000Z"), ed_opp.strftime("%Y-%m-%dT23:59:59.999Z")
            s_u_c, e_u_c = make_utc_range(sd_con, ed_con)
            fb_s, fb_h = self.fb_picker.start_date.strftime("%Y-%m-%d"), self.fb_picker.end_date.strftime("%Y-%m-%d")
            self.log("Extrayendo..."); res_o, res_v, res_c, res_fb = [], [], [], []
            me_pages = obtener_paginas_autorizadas()
            with ThreadPoolExecutor(max_workers=5) as ex:
                f_opp = {ex.submit(fetch_for_account, acc, ghl_s_o, ghl_e_o, s_iso_o, e_iso_o, self.log): acc for acc in ACCOUNTS}
                f_con = {ex.submit(fetch_contacts_for_account, acc, s_u_c, e_u_c, self.log): acc for acc in ACCOUNTS}
                f_fb = [ex.submit(obtener_insights, acc, fb_s, fb_h, self.log) for acc in FB_AD_ACCOUNTS]
                for f in as_completed(list(f_opp.keys()) + list(f_con.keys()) + f_fb):
                    if f in f_opp: o, v = f.result(); res_o.extend(o); res_v.extend(v)
                    elif f in f_con: c_data = f.result(); self.log(f"  {f_con[f]['name']}: {len(c_data)} contactos."); res_c.extend(c_data)
                    else:
                        insights = f.result()
                        if insights:
                            self.log(f"  Facebook: {len(insights)} líneas de gasto extraídas.")
                            ad_ids = list({i["ad_id"] for i in insights if "ad_id" in i})
                            creative_map = obtener_creatives(ad_ids)
                            p_map = obtener_paginas(list(set(creative_map.values())))
                            all_page_names = {**me_pages, **obtener_nombres_paginas(list(set(p_map.values())))}
                            for ins in insights:
                                aid, ad_n, camp = str(ins.get("ad_id")), ins.get("ad_name", ""), ins.get("campaign_name", "")
                                acc_id = str(ins.get("account_id"))
                                pid = p_map.get(creative_map.get(aid)); pname = all_page_names.get(pid)
                                conv = int(next((a["value"] for a in ins.get("actions", []) if a["action_type"] == "onsite_conversion.messaging_conversation_started_7d"), 0))
                                anu, tpost = extraer_datos_anuncio(ad_n)
                                spend = float(ins.get("spend", 0))
                                if f"act_{acc_id}" in FB_USD_ACCOUNTS or acc_id in FB_USD_ACCOUNTS:
                                    spend *= USD_TO_GTQ

                                d = ins.get("date_start", "")
                                res_fb.append({
                                    "ID del anuncio": aid,
                                    "ID de la página": pid,
                                    "Nombre de la página": pname,
                                    "Nombre de la campaña": camp,
                                    "Nombre del conjunto": ins.get("adset_name"),
                                    "Nombre del anuncio": ad_n,
                                    "Día": d,
                                    "Contactos mensajes nuevos": conv,
                                    "Importe gastado": spend,
                                    "Mes": int(d[5:7]) if d else 0,
                                    "Anio": int(d[0:4]) if d else 0,
                                    "codigo": anu,
                                    "precio": extraer_precio_fb(ad_n),
                                    "tipo_post": tpost,
                                    "SECUENCIA": mapping_secuencia_gasto(camp)
                                })

            df_metas = cargar_metas(PATH_METAS, self.log)



            if res_o or res_c or res_fb:
                self.log(f"Total extraído: {len(res_o)} ventas, {len(res_c)} contactos, {len(res_fb)} líneas de gasto.")

                # 1. Mapa de cruce: ID Contacto -> (Anuncio, Secuencia) para enriquecer ventas
                c_map = {c['id']: (c.get('anuncio',''), c.get('secuencia','')) for c in res_c if c.get('id')}
                for o in res_o:
                    cid = o.get('ID de contacto')
                    if cid in c_map:
                        if not o.get('Anuncio') or str(o.get('Anuncio','')).strip() == "":
                            o['Anuncio'] = c_map[cid][0]
                        if not o.get('Secuencia') or str(o.get('Secuencia','')).strip() == "":
                            o['Secuencia'] = c_map[cid][1]

                # 2. Asegurar campos Mes/Año en todas las fuentes
                for c in res_c:
                    d = c.get('fecha_iso','')
                    if 'Mes' not in c: c['Mes'] = int(d[5:7]) if d else 0
                    if 'Anio' not in c: c['Anio'] = int(d[0:4]) if d else 0

                # 3. Procesar costos de contactos
                df_c_final = self.process_contact_costs(res_c, res_fb)

                # 4. Cruzar VENTAS -> CONTACTOS (Saber cuánto vendió cada contacto/anuncio)
                v_map = {} # ID Contacto -> Total Venta
                for o in res_o:
                    cid = o.get('ID de contacto')
                    if cid:
                        v_map[cid] = v_map.get(cid, 0) + float(o.get('Valor del cliente potencial', 0))

                # Inyectar valor de venta en el DataFrame de contactos final
                if not df_c_final.empty:
                    df_c_final['Valor Venta'] = df_c_final['id'].map(v_map).fillna(0.0)

                # 5. Cruzar COSTOS -> VENTAS (Saber cuánto costó cada venta basado en su lead)
                c_cost_map = df_c_final.set_index('id')['Costo Total'].to_dict() if not df_c_final.empty else {}
                for o in res_o:
                    cid = o.get('ID de contacto')
                    o['Costo Lead'] = c_cost_map.get(cid, 0.0)

                self.generate_excel(res_o, res_v, df_c_final.to_dict(orient="records"), res_fb)
                self.generate_dashboard_html(res_o, res_v, df_c_final.to_dict(orient="records"), res_fb, df_metas)

            else: self.log("Sin datos.")
        except Exception as e: err_msg = traceback.format_exc(); print(err_msg); self.log(f"Error: {str(e)}\nConsulte la consola para detalles.")
        finally: self.after(0, lambda: self.generate_btn.configure(state="normal", text="🚀 GENERAR EXCEL"))

    def generate_dashboard_html(self, res_o, res_v, res_c, res_fb, df_metas):
        self.log("Generando Dashboard HTML...")
        import json
        import pandas as pd
        from datetime import datetime

        def prepare_json(data):
            try:
                if data is None: return []
                df = pd.DataFrame(data) if isinstance(data, list) else data.copy()
                if df.empty: return []
                # Limpiar infinitos antes de convertir a JSON
                df.replace([float('inf'), float('-inf')], 0, inplace=True)
                def clean_name(c):
                    s = str(c).lower().strip().replace(" ", "_")
                    s = s.replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u").replace("ñ", "n")
                    return "".join(ch for ch in s if ch.isalnum() or ch == "_")
                df.columns = [clean_name(c) for c in df.columns]
                df = df.loc[:, ~df.columns.duplicated(keep='first')]
                for col in df.columns:
                    if pd.api.types.is_numeric_dtype(df[col]):
                        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
                    elif pd.api.types.is_datetime64_any_dtype(df[col]):
                        df[col] = df[col].apply(lambda x: x.isoformat() if hasattr(x, "isoformat") else str(x))
                    else:
                        df[col] = df[col].fillna("").astype(str)
                return df.to_dict(orient="records")
            except Exception as e:
                print(f"Error en prepare_json: {e}")
                return []

        payload = {
            "oportunidades": prepare_json(res_o),
            "facebook": prepare_json(res_fb),
            "metas": prepare_json(df_metas),
            "contactos": prepare_json(res_c)
        }

        html_content = r"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>DUPAZA PRO - Dashboard</title>
    <script src="https://cdn.plot.ly/plotly-2.32.0.min.js"></script>
    <script src="https://cdn.tailwindcss.com"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;800&display=swap');
        body { font-family: 'Inter', sans-serif; background-color: #f8fafc; }
        .card { background: white; border-radius: 16px; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1); padding: 24px; border: 1px solid #e2e8f0; }
        .kpi-val { font-size: 28px; font-weight: 800; color: #0f172a; }
        .kpi-label { font-size: 11px; font-weight: 700; color: #64748b; text-transform: uppercase; letter-spacing: 0.05em; }
    </style>
</head>
<body class="p-6">
    <div class="max-w-7xl mx-auto">
        <header class="flex justify-between items-center mb-8 border-b pb-6">
            <div><h1 class="text-3xl font-black text-slate-800">📊 DUPAZA DASHBOARD</h1></div>
            <div class="text-right text-[10px] text-slate-400 font-bold uppercase">TIMESTAMP_HERE</div>
        </header>

        <div class="grid grid-cols-1 md:grid-cols-7 gap-4 mb-8 bg-white p-4 rounded-2xl shadow-sm border">
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">FECHA INICIO</label><input type="date" id="f-start" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">FECHA FIN</label><input type="date" id="f-end" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">GERENTE</label><select id="f-ger" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"><option value="ALL">TODOS</option></select></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">MARCA</label><select id="f-mar" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"><option value="ALL">TODAS</option></select></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">MES</label><select id="f-mes" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"><option value="ALL">TODOS</option></select></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">ASESOR</label><select id="f-ven" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"><option value="ALL">TODOS</option></select></div>
            <div><label class="block text-[10px] font-black text-slate-400 mb-1">ANUNCIO</label><select id="f-anu" class="w-full border rounded-lg p-2 text-xs font-bold outline-none focus:ring-2 focus:ring-green-500"><option value="ALL">TODOS</option></select></div>
        </div>

        <div class="grid grid-cols-2 md:grid-cols-5 gap-4 mb-8">
            <div class="card text-center"><div class="kpi-label">Gasto</div><div id="kpi-gasto" class="kpi-val text-blue-600">Q 0</div></div>
            <div class="card text-center"><div class="kpi-label">Leads</div><div id="kpi-leads" class="kpi-val">0</div></div>
            <div class="card text-center"><div class="kpi-label">Ventas</div><div id="kpi-venta" class="kpi-val text-emerald-600">Q 0</div></div>
            <div class="card text-center"><div class="kpi-label">ROAS</div><div id="kpi-roas" class="kpi-val">0.0</div></div>
            <div class="card text-center"><div class="kpi-label">% Meta</div><div id="kpi-meta" class="kpi-val">0%</div></div>
        </div>

        <div class="grid grid-cols-1 lg:grid-cols-2 gap-8 mb-8">
            <div class="card h-[400px]"><h3 class="kpi-label border-b pb-2 mb-4">Tendencia Diaria (Leads y Gasto)</h3><div id="ch-leads" class="h-full"></div></div>
            <div class="card h-[400px]"><h3 class="kpi-label border-b pb-2 mb-4">Ventas vs Meta</h3><div id="ch-meta" class="h-full"></div></div>
        </div>

        <div class="grid grid-cols-1 lg:grid-cols-2 gap-8 mb-8">
            <div class="card h-[400px]"><h3 class="kpi-label border-b pb-2 mb-4">Ventas por Marca</h3><div id="ch-marca" class="h-full"></div></div>
            <div class="card h-[400px]"><h3 class="kpi-label border-b pb-2 mb-4">Eficiencia por Asesor</h3><div id="ch-asesor" class="h-full"></div></div>
        </div>

        <div class="grid grid-cols-1 gap-8 mb-8">
            <div class="card h-[500px]"><h3 class="kpi-label border-b pb-2 mb-4">Rendimiento Detallado por Anuncio (Leads, Gasto y Ventas)</h3><div id="ch-ad-perf" class="h-full"></div></div>
        </div>

        <div id="debug" class="text-[10px] text-slate-400 font-mono bg-slate-100 p-4 rounded-xl">Cargando monitor...</div>
    </div>

    <script id="data" type="application/json">PAYLOAD_JSON</script>
    <script>
        (function() {
            let raw = null;
            try {
                raw = JSON.parse(document.getElementById('data').textContent);
                console.log("Datos cargados:", raw);
            } catch(e) {
                console.error("Error parseando JSON:", e);
                document.getElementById('debug').innerText = "ERROR: No se pudo cargar el JSON.";
                return;
            }

            function init() {
                try {
                    const dates = raw.facebook.map(f => f.dia).concat(raw.contactos.map(c => c.fecha_iso)).filter(Boolean).sort();
                    if (dates.length) {
                        document.getElementById("f-start").value = dates[0];
                        document.getElementById("f-end").value = dates[dates.length - 1];
                    }
                    const pop = (id, list) => {
                        const el = document.getElementById(id);
                        if (!el) return;
                        [...new Set(list)].filter(Boolean).sort().forEach(i => {
                            const o = document.createElement("option"); o.value = i; o.textContent = i; el.appendChild(o);
                        });
                    };
                    pop("f-ger", raw.metas.map(m => m.gerente || m.gerente_regional));
                    pop("f-mar", raw.metas.map(m => m.marca || m.empresa));
                    pop("f-ven", [...raw.contactos.map(c => c.asignado), ...raw.oportunidades.map(o => o.asignado)]);
                    pop("f-mes", raw.oportunidades.map(o => o.mes));
                    pop("f-anu", [...raw.contactos.map(c => c.anuncio), ...raw.facebook.map(f => f.codigo)]);

                    document.querySelectorAll("select, input[type='date']").forEach(s => s.onchange = update);
                    update();
                } catch(e) {
                    console.error("Error en init:", e);
                    document.getElementById('debug').innerText = "ERROR en init: " + e.message;
                }
            }

            function update() {
                try {
                    const start = document.getElementById("f-start").value;
                    const end = document.getElementById("f-end").value;
                    const g = document.getElementById("f-ger").value.toUpperCase();
                    const m = document.getElementById("f-mar").value.toUpperCase();
                    const v = document.getElementById("f-ven").value.toUpperCase();
                    const mes = document.getElementById("f-mes").value;
                    const anu = document.getElementById("f-anu").value.toUpperCase();

                    const seqMap = raw.metas.reduce((acc, c) => {
                        const s = (c.sub_anillo || c.secuencia || "").toUpperCase().trim();
                        if (!s) return acc;
                        if (!acc[s]) acc[s] = { gers: new Set(), marcs: new Set() };
                        const gVal = c.gerente || c.gerente_regional || "";
                        const mVal = c.marca || c.empresa || "";
                        if (gVal) acc[s].gers.add(gVal.toUpperCase());
                        if (mVal) acc[s].marcs.add(mVal.toUpperCase());
                        return acc;
                    }, {});

                    const f_c = raw.contactos.filter(c => {
                        const s = (c.secuencia || "").toUpperCase().trim();
                        const d = c.fecha_iso;
                        const mDate = (!start || d >= start) && (!end || d <= end);
                        const mG = (g === "ALL" || (seqMap[s] && seqMap[s].gers.has(g)));
                        const mM = (m === "ALL" || (seqMap[s] && seqMap[s].marcs.has(m)));
                        const mV = (v === "ALL" || (c.asignado || "").trim().toUpperCase() === v);
                        const mA = (anu === "ALL" || (c.anuncio || "").toUpperCase() === anu);
                        return mDate && mG && mM && mV && mA;
                    });

                    const f_o = raw.oportunidades.filter(o => {
                        const s = (o.secuencia || "").toUpperCase().trim();
                        const d = o.fecha_iso;
                        const mDate = (!start || d >= start) && (!end || d <= end);
                        const mG = (g === "ALL" || (seqMap[s] && seqMap[s].gers.has(g)));
                        const mM = (m === "ALL" || (o.marca || "").toUpperCase() === m || (seqMap[s] && seqMap[s].marcs.has(m)));
                        const mV = (v === "ALL" || (o.asignado || "").trim().toUpperCase() === v);
                        const mMes = (mes === "ALL" || String(o.mes) === String(mes));
                        const mA = (anu === "ALL" || (o.anuncio || "").toUpperCase() === anu);
                        return mDate && mG && mM && mV && mMes && mA;
                    });

                    const f_fb = raw.facebook.filter(f => {
                        const s = (f.secuencia || "").toUpperCase().trim();
                        const d = f.dia;
                        const mDate = (!start || d >= start) && (!end || d <= end);
                        const mG = (g === "ALL" || (seqMap[s] && seqMap[s].gers.has(g)));
                        const mM = (m === "ALL" || (seqMap[s] && seqMap[s].marcs.has(m)));
                        const mA = (anu === "ALL" || (f.codigo || "").toUpperCase() === anu);
                        return mDate && mG && mM && mA;
                    });

                    const tGto = (v !== "ALL" || anu !== "ALL")
                        ? f_c.reduce((a, c) => a + Number(c.costo_total || 0), 0)
                        : f_fb.reduce((a, f) => a + Number(f.importe_gastado || 0), 0);

                    const tVta = f_o.reduce((a, c) => a + Number(c.valor_del_cliente_potencial || 0), 0);
                    const tLds = f_c.length;

                    document.getElementById("kpi-gasto").innerText = "Q" + Math.round(tGto).toLocaleString();
                    document.getElementById("kpi-leads").innerText = tLds.toLocaleString();
                    document.getElementById("kpi-venta").innerText = "Q" + Math.round(tVta).toLocaleString();
                    document.getElementById("kpi-roas").innerText = tGto > 0 ? (tVta / tGto).toFixed(1) : "0.0";

                    render(f_o, f_c, f_fb, tVta, tGto);
                    const dbg = `Opps: ${raw.oportunidades.length} (filt: ${f_o.length}) | FB: ${raw.facebook.length} (filt: ${f_fb.length}) | Leads: ${raw.contactos.length} (filt: ${f_c.length}) | Metas: ${raw.metas.length} | seqMapKeys: ${Object.keys(seqMap).length}`;
                    console.log(dbg);
                    document.getElementById('debug').innerText = dbg;
                } catch(e) {
                    console.error("Error en update:", e);
                    document.getElementById('debug').innerText = "ERROR en update: " + e.message;
                }
            }

            function render(fo, fc, ffb, tv, tg) {
                const layout = { autosize: true, margin: {t:30, b:60, l:50, r:50}, paper_bgcolor: "rgba(0,0,0,0)", plot_bgcolor: "rgba(0,0,0,0)", font: {size: 10} };
                const config = { responsive: true };

                const daily = {};
                fc.forEach(c => { const d = c.fecha_iso || "N/A"; if(!daily[d]) daily[d] = { leads: 0, spend: 0 }; daily[d].leads += 1; });
                ffb.forEach(f => { const d = f.dia || "N/A"; if(!daily[d]) daily[d] = { leads: 0, spend: 0 }; daily[d].spend += Number(f.importe_gastado || 0); });
                const sortedDays = Object.keys(daily).sort();

                Plotly.newPlot("ch-leads", [
                    { x: sortedDays, y: sortedDays.map(d => daily[d].leads), name: "Leads", type: "bar", marker: {color: "#3b82f6", opacity: 0.7} },
                    { x: sortedDays, y: sortedDays.map(d => daily[d].spend), name: "Gasto (Q)", type: "scatter", mode: "lines+markers", yaxis: "y2", line: {color: "#ef4444", width:2} }
                ], { ...layout, yaxis: { title: "Leads" }, yaxis2: { title: "Gasto (Q)", overlaying: "y", side: "right" }, showlegend: true, legend: { orientation: "h", y: -0.2 } }, config);

                const vbm = fo.reduce((acc, c) => { const m = (c.marca || "OTRA").toUpperCase(); acc[m] = (acc[m] || 0) + Number(c.valor_del_cliente_potencial || 0); return acc; }, {});
                const sortedM = Object.entries(vbm).sort((a,b) => b[1] - a[1]);
                Plotly.newPlot("ch-marca", [{ labels: sortedM.map(x => x[0]), values: sortedM.map(x => x[1]), type: "pie", hole: .4, marker: { colors: ["#3b82f6", "#10b981", "#f59e0b", "#ef4444", "#8b5cf6"] } }], { ...layout, showlegend: true }, config);

                const vm = fo.reduce((acc, c) => { const n = c.asignado || "Sin Asignar"; acc[n] = (acc[n] || 0) + Number(c.valor_del_cliente_potencial || 0); return acc; }, {});
                const vs = Object.entries(vm).sort((a,b) => a[1] - b[1]);
                Plotly.newPlot("ch-asesor", [{ y: vs.map(x => x[0]), x: vs.map(x => x[1]), type: "bar", orientation: "h", marker: {color: "#10b981"} }], { ...layout, margin: {t:10, b:40, l:150, r:10} }, config);

                const activeG = document.getElementById("f-ger").value.toUpperCase();
                const activeM = document.getElementById("f-mar").value.toUpperCase();
                const activeMes = document.getElementById("f-mes").value;
                const metasF = raw.metas.filter(x => {
                    const mG = (activeG === "ALL" || (x.gerente || "").toUpperCase() === activeG);
                    const mM = (activeM === "ALL" || (x.marca || "").toUpperCase() === activeM);
                    const mMes = (activeMes === "ALL" || String(x.mes) === String(activeMes));
                    return mG && mM && mMes;
                });
                const tMeta = metasF.reduce((a, c) => a + Number(c.metas_valor || c.meta || c.metas || 0), 0);
                const perc = tMeta > 0 ? (tv / tMeta) * 100 : 0;
                document.getElementById("kpi-meta").innerText = Math.round(perc) + "%";
                Plotly.newPlot("ch-meta", [{ domain: { x: [0, 1], y: [0, 1] }, value: tv, title: { text: "Cumplimiento de Meta" }, type: "indicator", mode: "gauge+number", gauge: { axis: { range: [0, Math.max(tMeta, tv * 1.2)] }, bar: { color: "#10b981" }, steps: [{ range: [0, tMeta], color: "#e2e8f0" }] } }], layout, config);

                const adData = {};
                ffb.forEach(f => {
                    const a = (f.codigo || "SIN CODA").toUpperCase();
                    if (!adData[a]) adData[a] = { leads: 0, spend: 0, sales: 0 };
                    adData[a].spend += Number(f.importe_gastado || 0);
                });
                fc.forEach(c => {
                    const a = (c.anuncio || "SIN CODA").toUpperCase();
                    if (!adData[a]) adData[a] = { leads: 0, spend: 0, sales: 0 };
                    adData[a].leads += 1;
                });
                fo.forEach(o => {
                    const a = (o.anuncio || "SIN CODA").toUpperCase();
                    if (!adData[a]) adData[a] = { leads: 0, spend: 0, sales: 0 };
                    adData[a].sales += Number(o.valor_del_cliente_potencial || 0);
                });

                const ads = Object.entries(adData).sort((a,b) => b[1].spend - a[1].spend).slice(0, 20);
                Plotly.newPlot("ch-ad-perf", [
                    { x: ads.map(x => x[0]), y: ads.map(x => x[1]).map(v => v.leads), name: "Leads", type: "bar", marker: {color: "#3b82f6"} },
                    { x: ads.map(x => x[0]), y: ads.map(x => x[1]).map(v => v.sales), name: "Ventas (Q)", type: "scatter", mode: "lines+markers", yaxis: "y2", line: {color: "#10b981", width: 3} },
                    { x: ads.map(x => x[0]), y: ads.map(x => x[1]).map(v => v.spend), name: "Gasto (Q)", type: "scatter", mode: "lines+markers", yaxis: "y2", line: {color: "#ef4444", width: 3} }
                ], {
                    ...layout, showlegend: true, legend: { orientation: "h", y: -0.2 }, yaxis: { title: "Cant. Leads" }, yaxis2: { title: "Dinero (Q)", overlaying: "y", side: "right" }
                }, config);
            }
            init();
        })();
    </script>
</body>
</html>"""
        final = html_content.replace("PAYLOAD_JSON", json.dumps(payload))
        final = final.replace("TIMESTAMP_HERE", datetime.now().strftime("%d/%m/%Y %H:%M"))
        with open("index.html", "w", encoding="utf-8") as f: f.write(final)
        self.log("Dashboard unificado listo.")




    def process_contact_costs(self, res_c, res_fb):
        import pandas as pd
        df_c = pd.DataFrame(res_c)
        df_fb = pd.DataFrame(res_fb)

        if df_c.empty:
            return df_c

        # 1. Normalizar columnas base
        for col in ['fecha_iso', 'secuencia', 'anuncio']:
            if col not in df_c.columns: df_c[col] = ""
            else: df_c[col] = df_c[col].astype(str).str.strip().str.upper()

        if not df_fb.empty:
            for col in ['Día', 'SECUENCIA', 'codigo']:
                if col not in df_fb.columns: df_fb[col] = ""
                else: df_fb[col] = df_fb[col].astype(str).str.strip().str.upper()

            for col in ['Importe gastado', 'Contactos mensajes nuevos']:
                if col in df_fb.columns:
                    df_fb[col] = pd.to_numeric(df_fb[col], errors='coerce').fillna(0.0)
                else:
                    df_fb[col] = 0.0
        else:
            df_fb = pd.DataFrame(columns=['Día', 'SECUENCIA', 'codigo', 'Importe gastado', 'Contactos mensajes nuevos'])

        # 2. Atribución Ad-hoc para contactos vacíos (RANKING TOP 3 con ROTACIÓN)
        def get_top_ads(fb_df):
            if fb_df.empty: return pd.DataFrame(columns=['Día', 'SECUENCIA', 'codigo', 'Contactos mensajes nuevos'])
            grouped = fb_df.groupby(['Día', 'SECUENCIA', 'codigo'])['Contactos mensajes nuevos'].sum().reset_index()
            grouped = grouped.sort_values(['Día', 'SECUENCIA', 'Contactos mensajes nuevos'], ascending=[True, True, False])
            top3 = grouped.groupby(['Día', 'SECUENCIA']).head(3)
            return top3

        top3_ads = get_top_ads(df_fb)
        rotation_counters = {}

        def assign_top_ad(row, top_df):
            curr_anu = str(row.get('anuncio', '')).strip()
            if curr_anu != "" and curr_anu != "NAN" and curr_anu != "NONE": return curr_anu

            f_iso, seq = row.get('fecha_iso', ''), row.get('secuencia', '')
            if not f_iso or not seq: return ""

            key = (f_iso, seq)
            match = top_df[(top_df['Día'] == f_iso) & (top_df['SECUENCIA'] == seq)]
            if not match.empty:
                idx = rotation_counters.get(key, 0) % len(match)
                rotation_counters[key] = idx + 1
                return str(match.iloc[idx]['codigo'])
            return ""

        df_c['anuncio'] = df_c.apply(lambda r: assign_top_ad(r, top3_ads), axis=1)

        # 3. Gasto Directo por Anuncio
        fb_grouped = df_fb.groupby(['Día', 'SECUENCIA', 'codigo'])['Importe gastado'].sum().reset_index()
        c_counts = df_c.groupby(['fecha_iso', 'secuencia', 'anuncio']).size().reset_index(name='contact_count')

        direct_costs = pd.merge(
            c_counts,
            fb_grouped,
            left_on=['fecha_iso', 'secuencia', 'anuncio'],
            right_on=['Día', 'SECUENCIA', 'codigo'],
            how='inner'
        )
        direct_costs['cost_per_contact'] = direct_costs['Importe gastado'] / direct_costs['contact_count'].replace(0, 1)

        df_c = pd.merge(
            df_c,
            direct_costs[['fecha_iso', 'secuencia', 'anuncio', 'cost_per_contact']],
            on=['fecha_iso', 'secuencia', 'anuncio'],
            how='left'
        )
        df_c['Costo Directo'] = pd.to_numeric(df_c['cost_per_contact'], errors='coerce').fillna(0.0)

        # 4. Gasto Repartido (Huérfanos)
        fb_attributed_keys = set(zip(direct_costs['fecha_iso'], direct_costs['secuencia'], direct_costs['anuncio']))
        df_fb['is_orphan'] = df_fb.apply(lambda r: (r['Día'], r['SECUENCIA'], r['codigo']) not in fb_attributed_keys, axis=1)

        orphan_spend = df_fb[df_fb['is_orphan']].groupby(['Día', 'SECUENCIA'])['Importe gastado'].sum().reset_index(name='total_orphan_spend')
        seq_total_contacts = df_c.groupby(['fecha_iso', 'secuencia']).size().reset_index(name='seq_total')

        allocation_base = pd.merge(orphan_spend, seq_total_contacts, left_on=['Día', 'SECUENCIA'], right_on=['fecha_iso', 'secuencia'])
        allocation_base['orphan_cost_per_contact'] = allocation_base['total_orphan_spend'] / allocation_base['seq_total'].replace(0, 1)

        df_c = pd.merge(
            df_c,
            allocation_base[['fecha_iso', 'secuencia', 'orphan_cost_per_contact']],
            on=['fecha_iso', 'secuencia'],
            how='left'
        )
        df_c['Gasto Repartido'] = pd.to_numeric(df_c['orphan_cost_per_contact'], errors='coerce').fillna(0.0)

        # 5. Costo Total
        df_c['Costo Total'] = df_c['Costo Directo'] + df_c['Gasto Repartido']

        drop_cols = ['cost_per_contact', 'orphan_cost_per_contact']
        df_c.drop(columns=[c for c in drop_cols if c in df_c.columns], inplace=True)
        return df_c

    def generate_excel(self, res_o, res_v, res_c, res_fb):
        self.log("Compilando..."); df_o, df_v, df_c, df_fb = pd.DataFrame(res_o), pd.DataFrame(res_v), pd.DataFrame(res_c), pd.DataFrame(res_fb)
        if not df_o.empty: df_o["Valor del cliente potencial"] = pd.to_numeric(df_o["Valor del cliente potencial"], errors="coerce").fillna(0)
        for df in [df_o, df_v, df_c, df_fb]:
            if not df.empty:
                for col in df.select_dtypes(include=['datetime64[ns, UTC]', 'datetime64[ns, America/Guatemala]']).columns: df[col] = df[col].dt.tz_localize(None)
        head = ["Asignado", "Secuencia", "Mes", "Anio", "Anuncio", "fase", "Valor del cliente potencial", "Costo Lead", "Creado", "Ultimo Actualizado", "Seguidores", "Notas", "etiquetas", "estado", "Fecha de Venta", "NIT", "Camas y Combos SKU", "Cantidad Camas y Combo SKU", "Camas y Combos SKU1", "Cantidad Camas y Combo SKU1", "Cocinas SKU", "Cantidad Cocinas SKU", "Cocinas SKU1", "Cantidad Cocinas SKU1", "Salas SKU", "Cantidad Salas SKU", "Salas SKU1", "Cantidad Salas SKU1"]
        tail = ["", "Departamento", "Municipio", "Telefono 1", "Telefono 2", "ID de oportunidad", "ID de contacto", "Cliente", "Cod", "DataVenta", "Fecha", "MARCA", "ANILLO", "UBICACION"]
        if not df_o.empty:
            if "" not in df_o.columns: df_o[""] = ""
            for c in head + tail:
                if c not in df_o.columns: df_o[c] = ""
            extra = [c for c in df_o.columns if c not in set(head + tail)]; df_o = df_o[head + extra + tail]
        v_cols = ["ID CONTACTO", "NIT", "NOMBRE", "TEL1", "TEL2", "VENDEDOR", "SECUENCIA", "MUNICIPIO", "DIRECCION", "RCF", "canal", "DEPARTAMENTO", "FECHA", "SKU", "DESCRIPCION", "Cantidad de combo", "MARCA", "UBICACION", "ANILLO", "COMENTARIOS", "ID Oportunidad", "BODEGAF", "TOTAL DOCTO", "PRECIO COMBO"]
        if not df_v.empty:
            for c in v_cols:
                if c not in df_v.columns: df_v[c] = ""
            df_v = df_v[v_cols]
        c_cols = ["id", "fecha", "asignado", "secuencia", "Anuncio", "tipo_post", "Mes", "Anio", "Costo Directo", "Gasto Repartido", "Costo Total", "Valor Venta"]
        if not df_c.empty:
            if "anuncio" in df_c.columns: df_c.rename(columns={"anuncio": "Anuncio"}, inplace=True)
            if "Mes" not in df_c.columns: df_c["Mes"] = ""
            if "Anio" not in df_c.columns: df_c["Anio"] = ""
            for c in c_cols:
                if c not in df_c.columns: df_c[c] = ""
            df_c = df_c[c_cols]
        fn = f"reporte_Dupaza_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        with pd.ExcelWriter(fn, engine='openpyxl') as writer:
            if not df_o.empty: df_o.to_excel(writer, sheet_name='REPORTE', index=False)
            if not df_v.empty: df_v.to_excel(writer, sheet_name='VENTAS', index=False)
            if not df_c.empty: df_c.to_excel(writer, sheet_name='CONTACTOS', index=False)
            if not df_fb.empty: df_fb.to_excel(writer, sheet_name='FACEBOOK ADS', index=False)
            if not df_o.empty:
                hoja1_ids = df_o[['ID de oportunidad']].astype(str).apply(lambda x: x.str.strip())
                hoja1_ids.to_excel(writer, sheet_name='Hoja1', index=False, header=False)
            else: pd.DataFrame().to_excel(writer, sheet_name='Hoja1', index=False)
            h_f = PatternFill(start_color="76933C", end_color="76933C", fill_type="solid")
            h_font = Font(bold=True, color="FFFFFF")
            h_align = Alignment(horizontal="center")
            if not df_v.empty:
                ws_v = writer.book['VENTAS']; idx_bus = len(v_cols) + 1; ws_v.cell(row=1, column=idx_bus).value = "BUSQUEDA"
                for cell in ws_v[1]: cell.fill, cell.font, cell.alignment = h_f, h_font, h_align
                bg_f, cur_f = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid"), None
                for r in range(2, ws_v.max_row + 1):
                    val_id = ws_v.cell(row=r, column=1).value
                    if val_id and str(val_id).strip(): ws_v.cell(row=r, column=idx_bus).value = f"=VLOOKUP(U{r},Hoja1!A:A,1,FALSE)"; cur_f = bg_f if cur_f is None else None
                    if cur_f:
                        for c in range(1, idx_bus + 1): ws_v.cell(row=r, column=c).fill = cur_f
            if not df_c.empty:
                ws_c = writer.book['CONTACTOS']
                for cell in ws_c[1]: cell.fill, cell.font, cell.alignment = h_f, h_font, h_align
                for r in range(2, ws_c.max_row + 1):
                    for c_idx in range(7, 12): ws_c.cell(row=r, column=c_idx).number_format = '"Q" #,##0.00'
            if not df_o.empty:
                ws_o = writer.book['REPORTE']
                for cell in ws_o[1]: cell.fill, cell.font, cell.alignment = h_f, h_font, h_align
        self.log(f"ÉXITO: {fn}"); messagebox.showinfo("ÉXITO", f"Excel generado:\n{fn}")

if __name__ == "__main__":
    app = App(); app.mainloop()
